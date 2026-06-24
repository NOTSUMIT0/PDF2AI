from fastapi import APIRouter, UploadFile, File

from app.services.markitdown_service import (
    convert_document_to_markdown
)

from app.services.pdf_analysis import (
    analyze_pdf
)

from app.services.audio_transcriber import (
    transcribe_audio
)

from app.services.video_transcriber import (
    transcribe_video
)

from pydantic import BaseModel

from app.services.youtube_transcriber import (
    transcribe_youtube
)

from docx import Document

from pptx import Presentation

import zipfile
import shutil

import tempfile
import os
import traceback

router = APIRouter()

class YouTubeRequest(
    BaseModel
):

    url: str


def analyze_pptx(path):

    prs = Presentation(path)

    return {

        "pages":
            len(prs.slides),

        "images": 0,

        "scanned": False,

        "ocr_required": False,

        "estimated_time":
            "5-10 sec"

    }

def analyze_docx(path):

    doc = Document(path)

    paragraphs = len(
        doc.paragraphs
    )

    pages = max(
        1,
        round(
            paragraphs / 25
        )
    )

    return {

        "pages": pages,

        "images": 0,

        "scanned": False,

        "ocr_required": False,

        "estimated_time":
            "5-10 sec"

    }

@router.post("/youtube")
async def youtube_convert(
    request: YouTubeRequest
):

    markdown = transcribe_youtube(
        request.url
    )

    return {

        "success": True,

        "documents": [

            {

                "name":
                    "YouTube Transcript",

                "markdown":
                    markdown

            }

        ]

    }


@router.post("/analyze")
async def analyze_document(
    file: UploadFile = File(...)
):

    temp_path = None

    try:

        extension = (
            os.path.splitext(
                file.filename
            )[1]
            .lower()
        )

        with tempfile.NamedTemporaryFile(
            delete=False,
            suffix = os.path.splitext(file.filename)[1]
        ) as temp_file:

            content = await file.read()

            temp_file.write(content)

            temp_path = temp_file.name

        if extension == ".pptx":

            return {

                "success": True,

                "analysis":
                    analyze_pptx(
                        temp_path
                    )

            }    

        if extension == ".docx":

            return {

                "success": True,

                "analysis":
                    analyze_docx(
                        temp_path
                    )

            } 
        
        if extension in [
            ".mp3",
            ".wav",
            ".m4a",
            ".aac",
            ".flac",
            ".ogg"
        ]:
            
            file_size_mb = round(
                os.path.getsize(temp_path)
                / 1024
                / 1024,
                2
            )

            if file_size_mb < 5:
                estimated = "5-15 sec"

            elif file_size_mb < 20:
                estimated = "15-60 sec"

            else:
                estimated = "1-3 min"

            return {

                "success": True,

                "analysis": {

                    "pages": "Audio",

                    "images": "N/A",

                    "scanned": False,

                    "ocr_required": False,

                    "estimated_time":
                        estimated

                }

            }
        
        if extension in [

            ".mp4",
            ".mov",
            ".mkv",
            ".avi",
            ".webm"

        ]:

            file_size_mb = round(
                os.path.getsize(temp_path)
                / 1024
                / 1024,
                2
            )

            if file_size_mb < 50:

                estimated = "15-60 sec"

            elif file_size_mb < 200:

                estimated = "1-3 min"

            else:

                estimated = "3-10 min"

            return {

                "success": True,

                "analysis": {

                    "pages": "Video",

                    "images": "N/A",

                    "scanned": False,

                    "ocr_required": False,

                    "estimated_time":
                        estimated

                }

            }

        analysis = analyze_pdf(
            temp_path
        )

        return {

            "success": True,

            "analysis": analysis

        }

    finally:

        if (
            temp_path and
            os.path.exists(temp_path)
        ):
            os.remove(temp_path)



@router.post("/convert")
async def convert_document(
    file: UploadFile = File(...)
):
    temp_path = None

    try:

        with tempfile.NamedTemporaryFile(
            delete=False,
            suffix=os.path.splitext(
                file.filename
            )[1]
        ) as temp_file:

            content = await file.read()

            temp_file.write(
                content
            )

            temp_path = (
                temp_file.name
            )

        extension = (
            os.path.splitext(
                file.filename
            )[1]
            .lower()
        )

        audio_extensions = [

            ".mp3",

            ".wav",

            ".m4a",

            ".flac",

            ".ogg"

        ]

        if extension in audio_extensions:

            markdown = (
                transcribe_audio(
                    temp_path
                )
            )

            return {

                "success": True,

                "documents": [

                    {

                        "name":
                            file.filename,

                        "markdown":
                            markdown

                    }

                ]

            }

        # ZIP Processing
        if extension == ".zip":

            extracted_documents = []

            extract_dir = tempfile.mkdtemp()

            try:

                with zipfile.ZipFile(
                    temp_path,
                    "r"
                ) as zip_ref:

                    zip_ref.extractall(
                        extract_dir
                    )

                supported_extensions = [

                    ".pdf",
                    ".doc",
                    ".docx",
                    ".ppt",
                    ".pptx",
                    ".txt",
                    ".md",
                    ".html",
                    ".csv",
                    ".json",
                    ".xml",
                    ".mp3",
                    ".wav",
                    ".m4a",
                    ".flac",
                    ".ogg",

                ]

                for root, dirs, files in os.walk(
                    extract_dir
                ):

                    for extracted_file in files:

                        file_path = os.path.join(
                            root,
                            extracted_file
                        )

                        file_ext = (
                            os.path.splitext(
                                extracted_file
                            )[1]
                            .lower()
                        )

                        if (
                            file_ext
                            not in supported_extensions
                        ):
                            continue

                        print(
                            f"Processing ZIP file: {extracted_file}"
                        )

                        markdown = (
                            convert_document_to_markdown(
                                file_path,
                                extracted_file
                            )
                        )

                        extracted_documents.append(

                            {

                                "name":
                                    extracted_file,

                                "markdown":
                                    markdown

                            }

                        )

                return {

                    "success": True,

                    "documents":
                        extracted_documents

                }

            finally:

                shutil.rmtree(
                    extract_dir,
                    ignore_errors=True     
                )

        if extension in [

            ".mp4",
            ".mov",
            ".mkv",
            ".avi",
            ".webm"

        ]:

            markdown = (
                transcribe_video(
                    temp_path
                )
            )

            return {

                "success": True,

                "documents": [

                    {

                        "name":
                            file.filename,

                        "markdown":
                            markdown

                    }

                ]

            }        

        # NORMAL FILE PROCESSING

        markdown = (
            convert_document_to_markdown(
                temp_path,
                file.filename
            )
        )

        return {

            "success": True,

            "documents": [

                {

                    "name":
                        file.filename,

                    "markdown":
                        markdown

                }

            ]

        }

    except Exception as e:

        print(
            "CONVERT ERROR"
        )

        traceback.print_exc()

        error_text = str(e)

        if "403" in error_text:

            error_text = (
                "YouTube blocked the request. "
                "Try again in a few minutes."
            )

        elif "ffmpeg" in error_text.lower():

            error_text = (
                "Video processing failed. "
                "FFmpeg is missing."
            )

        elif "whisper" in error_text.lower():

            error_text = (
                "Speech recognition failed."
            )

        return {

            "success": False,

            "error":
                error_text

        }

    finally:

        if (
            temp_path and
            os.path.exists(
                temp_path
            )
        ):

            os.remove(
                temp_path
            )